from io import BytesIO
import json

from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_inventory_api import _client
from tests.test_panel_operations import wait_status


def test_inventory_export_owned_operation_download_and_revoked_rights(tmp_path,monkeypatch):
    monkeypatch.setenv('PANEL_EXPORT_ROOT',str(tmp_path/'exports'))
    client,headers = _client(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    from app.models import WebUser,PanelOperation,DownloadToken
    with get_sessionmaker()() as db:
        user = db.scalar(select(WebUser))
        user.is_admin = False
        user.permissions_json = '["inventory:export"]'
        db.commit()
    response = client.post('/inventory/export',data={'csrf_token':headers['X-CSRF-Token'],'scope':'all'},follow_redirects=False)
    assert response.status_code == 303,response.text
    operation = wait_status(response.headers['x-operation-id'],'succeeded')
    ids = json.loads(operation.artifact_ids_json)
    assert len(ids) == 1
    link = f'/operations/{operation.id}/files/{ids[0]}'
    result = client.get(link)
    assert result.status_code == 200
    assert load_workbook(BytesIO(result.content))['Устройства'].max_row == 1
    with get_sessionmaker()() as db:
        record = db.get(DownloadToken,ids[0])
        assert record.file_type == 'inventory-xlsx'
        db.scalar(select(WebUser)).permissions_json = '[]'
        db.commit()
    assert client.get(link).status_code == 403
    assert client.post('/inventory/export',data={'csrf_token':headers['X-CSRF-Token'],'scope':'all'}).status_code == 403
    with get_sessionmaker()() as db:
        db.scalar(select(WebUser)).permissions_json = '["inventory:export"]'
        db.get(PanelOperation,operation.id).owner = 'user:other'
        db.commit()
    assert client.get(link).status_code == 404


def test_deleted_export_requires_delete_permission_and_has_no_artifact_on_denial(tmp_path,monkeypatch):
    monkeypatch.setenv('PANEL_EXPORT_ROOT',str(tmp_path/'exports'))
    client,headers = _client(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    from app.models import WebUser
    with get_sessionmaker()() as db:
        user = db.scalar(select(WebUser))
        user.is_admin = False
        user.permissions_json = '["inventory:export"]'
        db.commit()
    response = client.post('/inventory/export',data={'csrf_token':headers['X-CSRF-Token'],'scope':'deleted'},follow_redirects=False)
    assert response.status_code == 303
    operation = wait_status(response.headers['x-operation-id'],'failed')
    assert json.loads(operation.artifact_ids_json) == []
    assert not (tmp_path/'exports').exists()
    assert client.get(f'/operations/{operation.id}/result').status_code == 403


def test_inventory_new_export_intent_reads_new_facts_and_transport_retry_reuses_result(tmp_path,monkeypatch):
    import re
    monkeypatch.setenv('PANEL_EXPORT_ROOT',str(tmp_path/'exports'))
    client,headers = _client(tmp_path,monkeypatch)
    def key():
        return re.search(r'name="operation_key" value="([0-9a-f]{32})"',client.get('/inventory').text).group(1)
    first_key = key()
    data = {'csrf_token':headers['X-CSRF-Token'],'scope':'all','operation_key':first_key}
    first = client.post('/inventory/export',data=data,follow_redirects=False)
    first_id = first.headers['x-operation-id']
    wait_status(first_id,'succeeded')
    from app.db import get_sessionmaker
    from app.inventory.models import InventoryAsset,InventoryAssetType
    with get_sessionmaker()() as db:
        db.add(InventoryAsset(asset_type=InventoryAssetType.PC,custom_name='New facts'))
        db.commit()
    retry = client.post('/inventory/export',data=data,follow_redirects=False)
    assert retry.headers['x-operation-id'] == first_id
    second_key = key()
    assert second_key != first_key
    second = client.post('/inventory/export',data={**data,'operation_key':second_key},follow_redirects=False)
    assert second.headers['x-operation-id'] != first_id
    operation = wait_status(second.headers['x-operation-id'],'succeeded')
    file_id = json.loads(operation.artifact_ids_json)[0]
    payload = client.get(f'/operations/{operation.id}/files/{file_id}').content
    assert load_workbook(BytesIO(payload))['Устройства'].max_row == 2


def test_export_artifact_expiry_revocation_and_raw_token_route(tmp_path,monkeypatch):
    from datetime import datetime,timezone,timedelta
    monkeypatch.setenv('PANEL_EXPORT_ROOT',str(tmp_path/'exports'))
    client,headers = _client(tmp_path,monkeypatch)
    response = client.post('/inventory/export',data={'csrf_token':headers['X-CSRF-Token'],'scope':'all'},follow_redirects=False)
    operation = wait_status(response.headers['x-operation-id'],'succeeded')
    file_id = json.loads(operation.artifact_ids_json)[0]
    from app.db import get_sessionmaker
    from app.models import DownloadToken
    from app.download_tokens import create_download_token
    link = f'/operations/{operation.id}/files/{file_id}'
    now = datetime.now(timezone.utc)
    with get_sessionmaker()() as db:
        record = db.get(DownloadToken,file_id)
        file_path,owner = record.file_path,record.created_by
        record.expires_at = now-timedelta(seconds=1)
        db.commit()
    assert client.get(link).status_code == 404
    with get_sessionmaker()() as db:
        record = db.get(DownloadToken,file_id)
        record.expires_at = now+timedelta(minutes=10)
        record.revoked_at = now
        db.commit()
    assert client.get(link).status_code == 404
    token,_ = create_download_token(client_name='synthetic',file_path=file_path,file_type='inventory-xlsx',created_by=owner,expires_at=now+timedelta(minutes=10))
    assert client.get('/download/'+token).status_code == 404
