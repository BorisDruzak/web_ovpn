from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_inventory_lifecycle import workplace


def test_inventory_export_keeps_full_selection_details_and_outside_relations(tmp_path,monkeypatch):
    _,_,data = workplace(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    from app.inventory.models import InventoryAsset,InventoryAssetType,InventoryAssetRelation,InventoryPCDetails
    from app.inventory.xlsx_export import inventory_workbook
    factory = get_sessionmaker()
    with factory() as db:
        pc = db.get(InventoryAsset,data['pc']['id'])
        pc.inventory_number = '00123'
        pc.description = '=1+1'
        db.get(InventoryPCDetails,pc.id).ram_gb = 0
        outside = InventoryAsset(asset_type=InventoryAssetType.PRINTER,custom_name='Outside')
        db.add(outside)
        db.flush()
        db.add(InventoryAssetRelation(parent_asset_id=pc.id,child_asset_id=outside.id,created_by='synthetic'))
        db.add_all([InventoryAsset(asset_type=InventoryAssetType.OTHER,location_id=pc.location_id) for _ in range(251)])
        db.commit()
    payload,counts = inventory_workbook(factory,location_id=data['pc']['location_id'])
    book = load_workbook(BytesIO(payload))
    assert counts['assets'] == 255 and book['Устройства'].max_row == 256
    rows = list(book['Устройства'].values)
    headers = rows[0]
    pcrow = next(row for row in rows[1:] if row[0] == data['pc']['id'])
    assert pcrow[headers.index('Инвентарный номер')] == '00123'
    assert pcrow[headers.index('RAM ГБ')] == 0
    assert pcrow[headers.index('Описание')] == '=1+1'
    relationrows = list(book['Связи рабочего места'].values)
    assert len(relationrows) == 5
    assert 'Состояние' in relationrows[0]
    assert all(row[relationrows[0].index('Состояние')] == 'Действует' for row in relationrows[1:])
    assert any('вне выборки' in str(value) for row in relationrows for value in row)
    assert {'Сетевые привязки','Параметры','Идентификаторы','Проверки','Фото'} <= set(book.sheetnames)


def test_inventory_export_deleted_is_explicit_and_does_not_expand_scope(tmp_path,monkeypatch):
    _,_,data = workplace(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    from app.inventory.lifecycle import soft_delete
    from app.inventory.models import InventoryAsset
    from app.inventory.xlsx_export import inventory_workbook
    factory = get_sessionmaker()
    with factory() as db:
        asset = db.get(InventoryAsset,data['pc']['id'])
        soft_delete(db,asset.id,expected_revision=asset.manual_revision,actor='synthetic',reason='Historical')
        db.commit()
    _,active = inventory_workbook(factory)
    payload,deleted = inventory_workbook(factory,deleted=True)
    assert active['assets'] == 3 and deleted['assets'] == 1
    book = load_workbook(BytesIO(payload))
    assert book['Связи рабочего места'].max_row == 4
    relationships = list(book['Связи рабочего места'].values)
    assert all(row[relationships[0].index('Состояние')] == 'Завершена' for row in relationships[1:])
    assert all(row[0] == data['pc']['id'] for row in list(book['Устройства'].values)[1:])


def test_inventory_export_empty_and_budget_are_explicit(tmp_path,monkeypatch):
    import pytest
    _,_,data = workplace(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    import app.inventory.xlsx_export as exporter
    from app.xlsx_export import ExportLimit
    payload,counts = exporter.inventory_workbook(get_sessionmaker(),location_id='absent')
    assert counts['assets'] == 0 and load_workbook(BytesIO(payload))['Устройства'].max_row == 1
    monkeypatch.setattr(exporter,'MAX_SOURCE_ROWS',2)
    with pytest.raises(ExportLimit):
        exporter.inventory_workbook(get_sessionmaker())


def test_inventory_export_does_not_mix_concurrent_detail_revision(tmp_path,monkeypatch):
    _,_,data = workplace(tmp_path,monkeypatch)
    from sqlalchemy import event
    from app.db import get_engine,get_sessionmaker
    from app.inventory.xlsx_export import inventory_workbook
    engine = get_engine()
    with engine.connect() as connection:
        connection.exec_driver_sql('PRAGMA journal_mode=WAL')
    changed = []
    def concurrent_write(conn,cursor,statement,parameters,context,executemany):
        if 'FROM inventory_pc_details' in statement and not changed:
            changed.append(True)
            with engine.begin() as writer:
                writer.exec_driver_sql('UPDATE inventory_pc_details SET ram_gb=32 WHERE asset_id=?',(data['pc']['id'],))
    event.listen(engine,'before_cursor_execute',concurrent_write)
    try:
        payload,_ = inventory_workbook(get_sessionmaker())
    finally:
        event.remove(engine,'before_cursor_execute',concurrent_write)
    rows = list(load_workbook(BytesIO(payload))['Устройства'].values)
    pc = next(row for row in rows[1:] if row[0] == data['pc']['id'])
    assert changed and pc[rows[0].index('RAM ГБ')] == 8
    with engine.connect() as connection:
        assert connection.exec_driver_sql('SELECT ram_gb FROM inventory_pc_details WHERE asset_id=?',(data['pc']['id'],)).scalar_one() == 32


def test_inventory_export_rejects_source_text_before_orm_materialization(tmp_path,monkeypatch):
    import pytest
    from sqlalchemy import event
    _,_,data = workplace(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    from app.inventory.models import InventoryAsset
    import app.inventory.xlsx_export as exporter
    from app.xlsx_export import ExportLimit
    with get_sessionmaker()() as db:
        db.get(InventoryAsset,data['pc']['id']).description = '\x00'+'x'*1000
        db.commit()
    monkeypatch.setattr(exporter,'MAX_SOURCE_BYTES',500)
    loaded = []
    def on_load(target,context):
        loaded.append(target.id)
    event.listen(InventoryAsset,'load',on_load)
    try:
        with pytest.raises(ExportLimit):
            exporter.inventory_workbook(get_sessionmaker())
    finally:
        event.remove(InventoryAsset,'load',on_load)
    assert loaded == []


def test_inventory_export_preserves_endpoint_network_facts_and_profile_freshness(tmp_path,monkeypatch):
    from datetime import datetime,timezone,timedelta
    import json
    _,_,data = workplace(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    from app.inventory.models import InventoryExternalBinding,InventoryExternalBindingStatus,InventoryEndpointState
    from app.inventory.xlsx_export import inventory_workbook
    now = datetime.now(timezone.utc)
    old = (now-timedelta(hours=2)).isoformat()
    with get_sessionmaker()() as db:
        binding = InventoryExternalBinding(asset_id=data['pc']['id'],source='endpoint_platform',external_id='synthetic-device',
            status=InventoryExternalBindingStatus.CONFIRMED,binding_method='manual',confidence=100,created_by='synthetic')
        db.add(binding)
        db.flush()
        db.add(InventoryEndpointState(binding_id=binding.id,asset_id=binding.asset_id,endpoint_device_id=binding.external_id,
            online=False,last_success_at=now,baseline_snapshot_id='baseline-old',network_snapshot_id='network-old',
            safe_context_json={'ip':'192.0.2.24','mac':'02:00:00:00:00:24','hostname':'synthetic-host','secret':'NEVER_EXPORT',
                'profile_status':{'network_v1':'available'},'profile_collected_at':{'network_v1':old},'profile_checked_at':{'network_v1':old}}))
        db.commit()
    payload,_ = inventory_workbook(get_sessionmaker())
    rows = list(load_workbook(BytesIO(payload))['Сетевые привязки'].values)
    row = rows[1]
    for field,value in [('IP','192.0.2.24'),('MAC','02:00:00:00:00:24'),('Hostname','synthetic-host'),('Baseline снимок','baseline-old')]:
        assert row[rows[0].index(field)] == value
    assert json.loads(row[rows[0].index('Свежесть профилей')])['network_v1']['status'] == 'stale'
    assert 'NEVER_EXPORT' not in str(rows)


def test_inventory_export_keeps_netctl_presence_freshness_and_full_binding_history(tmp_path,monkeypatch):
    from datetime import datetime,timezone,timedelta
    import json
    _,_,data = workplace(tmp_path,monkeypatch)
    from app.db import get_sessionmaker
    from app.inventory.models import InventoryNetctlBinding,InventoryIdentifierSyncRun,InventoryExternalBindingStatus
    from app.inventory.xlsx_export import inventory_workbook
    now = datetime.now(timezone.utc)
    with get_sessionmaker()() as db:
        db.add(InventoryIdentifierSyncRun(snapshot_id=2,snapshot_generated_at=now,status='success',started_at=now,finished_at=now))
        for number in range(103):
            db.add(InventoryNetctlBinding(asset_id=data['pc']['id'],network_key=f'mac:02:00:00:00:01:{number:02X}',
                status=InventoryExternalBindingStatus.CONFIRMED,created_by='synthetic',observed_snapshot_id=1,observed_at=now,
                observation_json={'ip':'192.0.2.24','last_seen_at':now.isoformat(),'sources':['synthetic-switch'],
                    'availability':{'state':'seen'},'secret':'NEVER_EXPORT_NETCTL'}))
        db.commit()
    def facts():
        payload,counts = inventory_workbook(get_sessionmaker())
        rows = list(load_workbook(BytesIO(payload))['Сетевые привязки'].values)
        assert counts['bindings'] == 103 and len(rows) == 104
        assert 'NEVER_EXPORT_NETCTL' not in str(rows)
        return dict(zip(rows[0],rows[1]))
    saved = facts()
    assert saved['Состояние источника'] == 'available'
    assert saved['Присутствие в снимке'] == 'missing'
    assert saved['Свежесть наблюдения'] == 'fresh'
    assert saved['Доступность наблюдения'] == 'seen'
    assert json.loads(saved['Источники наблюдения']) == ['synthetic-switch']
    with get_sessionmaker()() as db:
        db.add(InventoryIdentifierSyncRun(status='failed',started_at=now+timedelta(seconds=1),
            finished_at=now+timedelta(seconds=1),failure_reason='synthetic unavailable'))
        db.commit()
    failed = facts()
    assert failed['Состояние источника'] == 'unavailable'
    assert failed['Присутствие в снимке'] == 'unknown'
    assert failed['IP'] == '192.0.2.24' and failed['Доступность наблюдения'] == 'seen'
