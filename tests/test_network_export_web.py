from io import BytesIO
import json
import subprocess
import sys

from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_inventory_api import _client
from tests.test_netctl_host_export import make_snapshot
from tests.test_panel_operations import wait_status


def test_full_network_export_real_cli_owned_download_and_filter(tmp_path,monkeypatch):
    monkeypatch.setenv('PANEL_EXPORT_ROOT',str(tmp_path/'exports'))
    client,headers = _client(tmp_path,monkeypatch)
    source,path = make_snapshot(tmp_path,300)
    source.close()
    before = path.read_bytes()
    calls = []
    def real_cli(args,**kwargs):
        calls.append(args)
        assert args[:2] == ['hosts','export'] and '--page' not in args and '--limit' not in args
        result = subprocess.run([sys.executable,'-m','netctl.cli','--json','--db',f'sqlite:///{path.as_posix()}',*args],
            input=kwargs.get('input_payload'),text=True,capture_output=True,timeout=20)
        assert result.returncode == 0,result.stdout+result.stderr
        return json.loads(result.stdout)
    monkeypatch.setattr('app.network_export.run_netctl',real_cli)
    from app.db import get_sessionmaker
    from app.models import WebUser
    with get_sessionmaker()() as db:
        user = db.scalar(select(WebUser))
        user.is_admin = False
        user.permissions_json = '["network:export"]'
        db.commit()
    response = client.post('/network/export',data={'csrf_token':headers['X-CSRF-Token'],'status':'all','seen_within':'all','inventory_link':'all'},follow_redirects=False)
    assert response.status_code == 303,response.text
    operation = wait_status(response.headers['x-operation-id'],'succeeded')
    file_id = json.loads(operation.artifact_ids_json)[0]
    link = f'/operations/{operation.id}/files/{file_id}'
    result = client.get(link)
    assert result.status_code == 200
    book = load_workbook(BytesIO(result.content))
    assert book['Снимок'].max_row == 301 and book['Связи с инвентаризацией'].max_row == 301
    assert path.read_bytes() == before and len(calls) == 1
    with get_sessionmaker()() as db:
        db.scalar(select(WebUser)).permissions_json = '[]'
        db.commit()
    assert client.get(link).status_code == 403


def test_network_export_rejects_concurrent_local_relation_change(tmp_path,monkeypatch):
    from fastapi import HTTPException
    import pytest
    _client(tmp_path,monkeypatch)
    source,path = make_snapshot(tmp_path,300)
    source.close()
    from app.db import get_sessionmaker
    from app.inventory.models import InventoryAsset,InventoryAssetType
    with get_sessionmaker()() as db:
        card = InventoryAsset(asset_type=InventoryAssetType.PC,custom_name='Before')
        db.add(card)
        db.commit()
        card_id = card.id
    def concurrent(args,**kwargs):
        result = subprocess.run([sys.executable,'-m','netctl.cli','--json','--db',f'sqlite:///{path.as_posix()}',*args],text=True,capture_output=True,timeout=20)
        with get_sessionmaker()() as db:
            db.get(InventoryAsset,card_id).custom_name = 'After'
            db.commit()
        return json.loads(result.stdout)
    monkeypatch.setattr('app.network_export.run_netctl',concurrent)
    from app.network_export import network_export_workbook
    with pytest.raises(HTTPException) as error:
        network_export_workbook({'status':'all','seen_within':'all','inventory_link':'all'})
    assert error.value.status_code == 409
