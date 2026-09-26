from datetime import datetime,timezone
from io import BytesIO
import pytest
from openpyxl import load_workbook


def test_network_workbook_complete_safe_sheets_and_metadata():
    from app.network_xlsx import network_workbook
    hosts=[{'ip':f'192.0.2.{n}','device_key':f'mac:02:00:00:00:01:{n:02X}','hostname':'=unsafe','mac':'00:01','status':'seen','sources':['arp'],
            'availability':{'state':'seen','passive_evidence':'arp','checked_at':'2026-09-26T10:00:00Z'},'secret':'DO_NOT_EXPORT'} for n in range(1,251)]
    enriched=[{**h,'inventory':{'state':'unlinked','asset':None},'endpoint_agent':{'state':'disabled','freshness':'unknown'}} for h in hosts]
    enriched[0]['inventory']={'state':'linked','binding_id':'relation','asset':{'id':'card','name':'Кириллица','inventory_number':'0001','manual_revision':4,'location':'Офис'}}
    result={'hosts':hosts,'total':250,'snapshot':{'snapshot_id':7,'generated_at':'2026-09-26T10:00:00Z','stale':True},'inventory_projection_revision':21}
    data,counts=network_workbook(result,enriched,filters={'status':'all','q':'secret-search','token':'SECRET'},inventory_epoch=21,enriched_at=datetime(2026,9,26,10,2,tzinfo=timezone.utc))
    book=load_workbook(BytesIO(data))
    assert book.sheetnames==['Снимок','Связи с инвентаризацией','Параметры']
    assert book['Снимок'].max_row==251 and book['Связи с инвентаризацией'].max_row==251
    assert counts=={'hosts':250,'linked':1,'unlinked':249}
    assert not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row)
    values=[cell.value for sheet in book for row in sheet for cell in row]
    assert 'DO_NOT_EXPORT' not in values and 'SECRET' not in values
    assert '0001' in values and 'Кириллица' in values and 'unlinked' in values
    assert all(s.freeze_panes=='A2' and s.auto_filter.ref for s in book)


def test_network_workbook_rejects_missing_or_reordered_enrichment():
    from app.network_xlsx import network_workbook
    result={'hosts':[{'ip':'1','device_key':'a'},{'ip':'2','device_key':'b'}],'total':2,'snapshot':{'snapshot_id':1}}
    for enriched in ([],list(reversed(result['hosts']))):
        with pytest.raises(ValueError):
            network_workbook(result,enriched,filters={},inventory_epoch=1,enriched_at=datetime.now(timezone.utc))


def test_network_workbook_accepts_real_snapshot_nested_passive_evidence(tmp_path):
    from netctl.db import connect
    from netctl.host_snapshot import refresh_host_snapshot,export_host_snapshot
    from app.network_xlsx import network_workbook
    conn=connect(f'sqlite:///{(tmp_path/"netctl.sqlite").as_posix()}')
    conn.execute("INSERT INTO network_hosts (ip,category,status,last_seen_at,tags_json) VALUES ('192.0.2.7','unknown','seen','2026-09-26T10:00:00Z','{}')");conn.commit()
    refresh_host_snapshot(conn,now='2026-09-26T10:00:00Z')
    result=export_host_snapshot(conn,{'status':'all'})
    data,counts=network_workbook(result,result['hosts'],filters={},inventory_epoch=1,enriched_at=datetime.now(timezone.utc))
    assert load_workbook(BytesIO(data))['Снимок'].max_row==2
    conn.close()


def test_network_workbook_empty_stale_snapshot_and_epoch_conflict():
    from app.network_xlsx import network_workbook
    result={'hosts':[],'total':0,'snapshot':{'snapshot_id':1,'stale':True},'inventory_projection_revision':7}
    with pytest.raises(ValueError,match='projection changed'):
        network_workbook(result,[],filters={},inventory_epoch=8,enriched_at=datetime.now(timezone.utc))
    data,counts=network_workbook(result,[],filters={},inventory_epoch=7,enriched_at=datetime.now(timezone.utc))
    book=load_workbook(BytesIO(data))
    assert book['Снимок'].max_row==1 and book['Связи с инвентаризацией'].max_row==1
    assert counts=={'hosts':0,'linked':0,'unlinked':0}
