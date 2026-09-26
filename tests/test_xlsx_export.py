from datetime import date
from io import BytesIO

from openpyxl import load_workbook


def test_workbook_never_creates_plaintext_worksheet_tempfiles(monkeypatch):
    from openpyxl.worksheet import _writer
    from app.xlsx_export import workbook_bytes
    def forbidden(*args, **kwargs):
        raise AssertionError('Worksheet plaintext tempfile attempted')
    monkeypatch.setattr(_writer,'create_temporary_file',forbidden)
    payload = workbook_bytes([('Устройства',['ID'],[['00123']])])
    assert load_workbook(BytesIO(payload)).active['A2'].value == '00123'


def test_workbook_rejects_invalid_xml_unicode_explicitly():
    import pytest
    from app.xlsx_export import workbook_bytes,ExportLimit
    for value in ['bad\ufffe','bad\uffff','bad\ud800']:
        with pytest.raises(ExportLimit):
            workbook_bytes([('Устройства',['ID'],[[value]])])


def test_workbook_rejects_duplicate_or_invalid_sheet_names_and_headers():
    import pytest
    from app.xlsx_export import workbook_bytes,ExportLimit
    for sheets in [
        [('Data',['ID'],[]),('data',['ID'],[])],
        [('bad/name',['ID'],[])],
        [('bad\ufffe',['ID'],[])],
        [('Data',[None],[])],
    ]:
        with pytest.raises(ExportLimit):
            workbook_bytes(sheets)


def test_workbook_preserves_types_and_formula_like_strings_as_safe_text():
    from app.xlsx_export import workbook_bytes
    dangerous = ['=1+1','+SUM(A1:A2)','-1','@SUM(A1:A2)',' \t=1+1','\x01=1+1','00123','Кириллица']
    rows = [[value,0,None,date(2026,9,26)] for value in dangerous]
    payload = workbook_bytes([('Устройства',['Номер','RAM','Неизвестно','Дата'],rows)])
    book = load_workbook(BytesIO(payload),data_only=False)
    sheet = book['Устройства']
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == 'A1:D9'
    for number,value in enumerate(dangerous,2):
        cell = sheet.cell(number,1)
        assert cell.data_type == 's'
        assert cell.value == value.replace('\x01','')
        assert sheet.cell(number,2).value == 0 and sheet.cell(number,2).data_type == 'n'
        assert sheet.cell(number,3).value is None
        assert sheet.cell(number,4).is_date
    assert not book._external_links and book.vba_archive is None


def test_workbook_rejects_oversized_cell_instead_of_truncating():
    import pytest
    from app.xlsx_export import workbook_bytes,ExportLimit
    with pytest.raises(ExportLimit):
        workbook_bytes([('Устройства',['Описание'],[['x'*32768]])])


def test_workbook_resource_budget_rejects_whole_result(monkeypatch):
    import pytest
    import app.xlsx_export as exporter
    monkeypatch.setattr(exporter,'MAX_ROWS',3)
    with pytest.raises(exporter.ExportLimit):
        exporter.workbook_bytes([('Устройства',['ID'],[[1],[2],[3]])])


def test_workbook_does_not_replace_empty_numeric_field_with_zero():
    from app.xlsx_export import workbook_bytes
    book = load_workbook(BytesIO(workbook_bytes([('Устройства',['RAM'],[[None],[0],[16]])])))
    assert [book['Устройства'].cell(row,1).value for row in range(2,5)] == [None,0,16]
